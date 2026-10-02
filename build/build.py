"""Rebuild the SCDOT AASHTOWare treatments workbook from the version the team is using.

cd build && python build.py ../source/AASHTOWare_Safety_Treatments_and_Countermeasures_input.xlsx \
    ../SCDOT_AASHTOWare_Safety_Treatments_and_Countermeasures_v8.xlsx
then recalculate in Excel or LibreOffice.
"""
import re
from copy import copy
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from criteria import CM, FIT, FIT_TEXT, FIT_DEFAULT_WHY

SRC, OUT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC)

F = "Arial"
MONO = "Consolas"
HDR_FILL = PatternFill("solid", fgColor="FF1A1A1A")
TITLE_FONT = Font(name=F, sz=14, b=True, color="FFFBD567")
HDR_FONT = Font(name=F, sz=9, b=True, color="FFFFFFFF")
BODY = Font(name=F, sz=9, color="FF222222")
BOLD = Font(name=F, sz=9, b=True, color="FF222222")
GREY = Font(name=F, sz=9, color="FF555555")
THIN = Side(style="thin", color="FFD0D0D0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(wrap_text=True, vertical="top")
FILL_OPEN = PatternFill("solid", fgColor="FFFFF4D6")
FILL_FIT = {
    "OK": PatternFill("solid", fgColor="FFE6F2E6"),
    "Understates": PatternFill("solid", fgColor="FFEFEFEF"),
    "Overstates": PatternFill("solid", fgColor="FFFCE4E4"),
    "Check": PatternFill("solid", fgColor="FFFFF4D6"),
}


def est_lines(text, width):
    if text is None:
        return 1
    n = 0
    for line in str(text).split("\n"):
        n += max(1, -(-len(line) // max(1, int(width * 1.15))))
    return n


def fit_height(ws, row, widths):
    lines = max(est_lines(ws.cell(row, c).value, w) for c, w in widths.items())
    ws.row_dimensions[row].height = max(15, min(409, 12.5 * lines + 4))


# ---------------------------------------------------------------- Countermeasures
cm = wb["Countermeasures"]
old_criteria = {r: cm.cell(r, 10).value for r in range(3, 33)}
names = {r: cm.cell(r, 2).value for r in range(3, 33)}
areas = {r: cm.cell(r, 1).value for r in range(3, 33)}

cm.auto_filter.ref = None
cm.freeze_panes = None
for rng in list(cm.merged_cells.ranges):
    cm.unmerge_cells(str(rng))
cm.delete_rows(1, cm.max_row)

HEAD = ["Emphasis area", "Countermeasure", "Area type", "Location", "Crash types",
        "Traffic control", "Site conditions",
        "Crash criteria (as entered in the SMS / AASHTOWare Condition Builder)",
        "Notes", "Open item", "# Treat."]
WID = [14, 30, 13, 13, 18, 10, 16, 60, 46, 34, 7]

cm.cell(1, 1, "COUNTERMEASURES").font = TITLE_FONT
cm.cell(1, 1).fill = HDR_FILL
for c in range(2, len(HEAD) + 1):
    cm.cell(1, c).fill = HDR_FILL
cm.row_dimensions[1].height = 25.5
cm.cell(2, 1, "Criteria use the field names and dropdown labels from the SMS Condition Builder. "
              "Each line is one row in the builder; keep the parentheses as groups. "
              "Codes and the rules behind them are on the Field Reference sheet.").font = GREY
cm.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEAD))
cm.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
cm.row_dimensions[2].height = 26
for j, h in enumerate(HEAD, 1):
    c = cm.cell(3, j, h)
    c.font = HDR_FONT
    c.fill = HDR_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
    cm.column_dimensions[c.column_letter].width = WID[j - 1]
cm.row_dimensions[3].height = 30

r_out = 4
cm_row_map = {}
for r in range(3, 33):
    atype, loc, ctypes, ctrl, site, crit, notes, open_item = CM[r]
    vals = [areas[r], names[r], atype, loc, ctypes, ctrl, site, crit, notes, open_item,
            f"=COUNTIF(Treatments!B:B,B{r_out})"]
    for j, v in enumerate(vals, 1):
        c = cm.cell(r_out, j, v)
        c.font = BODY
        c.alignment = TOP
        c.border = BOX
    cm.cell(r_out, 2).font = BOLD
    cm.cell(r_out, 8).font = Font(name=MONO, sz=9, color="FF1A1A1A")
    cm.cell(r_out, 11).alignment = Alignment(horizontal="center", vertical="top")
    if open_item:
        cm.cell(r_out, 10).fill = FILL_OPEN
    fit_height(cm, r_out, {8: 60 * 0.95, 9: 46, 10: 34, 2: 30})
    cm_row_map[r] = r_out
    r_out += 1

last_cm = r_out - 1
cm.freeze_panes = "C4"
cm.auto_filter.ref = f"A3:K{last_cm}"
cm.sheet_view.zoomScale = 90

# ---------------------------------------------------------------- Treatments
tr = wb["Treatments"]
SEV = [("K", 3), ("A", 4), ("B", 5), ("C", 6), ("O", 7)]
sev_changed = []
for r in range(3, 136):
    sev_text = str(tr.cell(r, 18).value or "")
    base = tr.cell(r, 3).value
    all_sev = sev_text.strip() in ("All", "Not specified", "")
    covered = {k for k, _ in SEV if all_sev or f"{k} (" in sev_text}
    if covered != {k for k, _ in SEV}:
        sev_changed.append(r)
    # P holds the CMF from the study; the five severity columns read from it.
    tr.cell(r, 16, base)
    tr.cell(r, 16).font = Font(name=F, sz=10, color="FF0000FF")
    for k, col in SEV:
        tr.cell(r, col, f'=IF(OR($R{r}="All",$R{r}="Not specified",ISNUMBER(SEARCH("{k} (",$R{r}))),$P{r},1)')
        tr.cell(r, col).font = Font(name=F, sz=10, color="FF000000")

tr.cell(2, 16, "CMF (study)")
tr.cell(2, 3).comment = Comment(
    "K-O columns read the study CMF in column P. A severity the study did not cover gets 1.00 "
    "(no effect), so AASHTOWare doesn't credit a reduction the study never measured.", "SCDOT")


def tidy_note(t):
    if not t:
        return t
    t = t.replace("Do NOT apply to collectors or local roads - the team ruled that out in review.",
                  "Not for collectors or locals (team).")
    t = t.replace("AADT range not reported by the study - the selector cannot score AADT fit for this row.",
                  "No AADT range in the study.")
    t = t.replace("AADT range not reported by the study", "No AADT range in the study.")
    t = re.sub(r"Study AADT [\d,\-]+ is well below SCDOT's typical volume on that functional class "
               r"\(~17,000 average, higher on the busiest routes\)\. Check the site AADT before using this CMF\.",
               "Study AADT is below the SC average for this class (~17,000).", t)
    t = t.replace("Study sites were ", "Study roads: ")
    t = t.replace("Valid AADT range tested:", "AADT tested:")
    return re.sub(r"\s+", " ", t).strip()


for r in range(3, 136):
    tr.cell(r, 33, tidy_note(tr.cell(r, 33).value))
tr.cell(2, 33, "Notes for SC use")

tr.cell(2, 34, "Fit with CM criteria")
tr.cell(2, 35, "Why")
for col in (34, 35):
    src = tr.cell(2, 33)
    dst = tr.cell(2, col)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
tr.column_dimensions["AH"].width = 13
tr.column_dimensions["AI"].width = 34
fit_counts = {}
for r in range(3, 136):
    f = FIT[r]
    key, why = (f, FIT_DEFAULT_WHY[f]) if isinstance(f, str) else f
    label = FIT_TEXT[key]
    if str(tr.cell(r, 30).value or "").startswith("No"):
        label, why = "Not used", "Excluded by the team"
    fit_counts[label] = fit_counts.get(label, 0) + 1
    a = tr.cell(r, 34, label)
    b = tr.cell(r, 35, why)
    for c in (a, b):
        c.font = Font(name=F, sz=10, color="FF222222")
        c.alignment = TOP
    if label in FILL_FIT:
        a.fill = FILL_FIT[label]
tr.cell(2, 34).comment = Comment(
    "OK: CMF measured the same crashes the criteria picks.\n"
    "Understates: CMF is for all crash types, applied here to a smaller group - conservative.\n"
    "Overstates: CMF measured a narrower group (night only, right-turn only...) than the criteria - "
    "benefit is too high.\nCheck: different crash type or facility.", "SCDOT")
tr.auto_filter.ref = "A2:AI135"
tr.freeze_panes = "C3"
for r in range(3, 136):
    tr.row_dimensions[r].height = None

# ---------------------------------------------------------------- Treatment Selector
sel = wb["Treatment Selector"]
for r in range(35, 168):
    tr_row = r - 32
    sel.cell(r, 27, f'=IF(OR(Treatments!AH{tr_row}="Overstates",Treatments!AH{tr_row}="Check"),'
                    f'UPPER(Treatments!AH{tr_row})&": "&Treatments!AI{tr_row}&". ","")'
                    f'&IF(Treatments!AG{tr_row}="","",Treatments!AG{tr_row})')
sel.column_dimensions["I"].width = 40

# ---------------------------------------------------------------- Summary
sm = wb["Summary"]
sm["A6"] = f"=COUNTA(Countermeasures!B4:B{last_cm})"
sm.cell(11, 1, "CMF fit with criteria").font = Font(name=F, sz=10, b=True)
rows = [("OK", "Same crashes"), ("Understates", "CMF is for all crash types"),
        ("Overstates", "CMF is for a narrower group"), ("Check", "Needs a decision")]
for i, (k, d) in enumerate(rows):
    sm.cell(12 + i, 1, k).font = Font(name=F, sz=10)
    sm.cell(12 + i, 2, f'=COUNTIF(Treatments!AH3:AH135,"{k}")').font = Font(name=F, sz=10)
    sm.cell(12 + i, 3, d).font = Font(name=F, sz=10, color="FF555555")
sm.cell(17, 1, "Severity-limited CMFs").font = Font(name=F, sz=10, b=True)
sm.cell(18, 1, "Treatments where K-O is not all the same").font = Font(name=F, sz=10)
sm.cell(18, 2, "=SUMPRODUCT(--((Treatments!C3:C135<>Treatments!G3:G135)+(Treatments!D3:D135<>Treatments!G3:G135)"
               "+(Treatments!E3:E135<>Treatments!G3:G135)+(Treatments!F3:F135<>Treatments!G3:G135)>0))").font = Font(name=F, sz=10)

# ---------------------------------------------------------------- Field Reference
fr = wb.create_sheet("Field Reference")
FR_HEAD = ["Field (as in the builder)", "Level", "Values used in this workbook (SMS code)",
           "New TR-310 / MMUCC equivalent", "Source / status"]
FR = [
 ("Junction Type", "Crash",
  "NON-JUNCTION (13), Cross-Over (01), Driveway (02), Four-Way Intersection (04), T-Intersection (08), "
  "Y-Intersection (12), Five/More Points (03), Traffic Circle (09)",
  "Split into Relation to Junction + Type of Intersection",
  "Confirmed: your Condition Builder query + SC crash data dictionary"),
 ("Manner of Collision", "Crash",
  "Non Collision (00), Rear End (10), Head On (20), Angle (41/42/43), Sideswipe, Same Direction (50), "
  "Sideswipe, Opposite Direction (60)",
  "Manner of Crash/Collision: Front to Rear, Front to Front, Angle, Sideswipe... New-form codes differ "
  "(Angle is 80 there, not in the SMS)",
  "Confirmed: your query + SC dictionary. Angle may list three times."),
 ("First Harmful Event / Crash Harmful Event", "Crash",
  "Cross Median/Center (02), Overturn/Rollover (08), Ran off Road Left (09), Ran off Road Right (10), "
  "Animal - Deer (20), Animal - Other (21), Pedalcycle (26), Pedestrian (27), Bridge Parapet End (41), "
  "Bridge Pier or Abutment (42), Bridge Rail (43), Culvert (44), Ditch (46), Embankment (47), "
  "Guardrail End (50), Guardrail Face (51), Median Barrier (56), Tree (60), Utility Pole (61)",
  "First Harmful Event. Under MMUCC, ran off road and cross median are NOT harmful events - "
  "they only go in Sequence of Events. New-form codes differ (Pedestrian 26, Pedalcycle 25).",
  "Confirmed: SC dictionary. Crash Harmful Event = most harmful event of the crash."),
 ("Sequence of Events", "Unit (one row per event)",
  "Same list as above. Used here: Cross Median/Center, Ran off Road Left",
  "Sequence of Events",
  "Use with '=' only. '≠' drops that row, not the crash."),
 ("First Harmful Event Location", "Crash",
  "Gore (1), Island (2), Median (3), Roadside (4), Roadway (5), Shoulder (6), Sidewalk (7), "
  "Outside Trafficway (8), Unknown (9)",
  "Location of First Harmful Event Relative to the Trafficway",
  "Confirmed: SMS dictionary you sent"),
 ("Trafficway", "Crash / Roadway",
  "Two-Way, Not Divided (1); Two-Way, Divided, Unprotected Median (2); Two-Way, Divided, Barrier (3); "
  "One-Way (4); Other (8)",
  "Trafficway Description (adds TWLTL as its own value)",
  "Confirmed: SMS dictionary + SC dictionary. Under Roadway, next to AADT."),
 ("Road Character", "Crash",
  "Curve-Level (4), Curve-On Grade (5), Curve-Hillcrest (6)",
  "Roadway Alignment and Grade (Curve Left / Curve Right)",
  "SC dictionary - confirm dropdown label"),
 ("Light Condition", "Crash",
  "Dark (Lighting Unspecified) (4), Dark (Street Lamp Lit) (5), Dark (Street Lamps Not Lit) (6), "
  "Dark (No Lights) (7)",
  "Light Condition: Dark-Lighted, Dark-Not Lighted, Dark-Unknown Lighting",
  "SC dictionary - confirm dropdown label"),
 ("Road Surface Condition", "Crash", "Wet (2)", "Roadway Surface Condition", "SC dictionary"),
 ("Traffic Control", "Crash", "Traffic Signal, Stop Sign", "Traffic Control Device Type",
  "Label not confirmed - check dropdown"),
 ("Probable Cause", "Unit",
  "Disregarded Sign/Signal (01), Exceeded Speed Limit (04), Obstruction in Roadway, Debris. "
  "Too Fast for Conditions (03) not used.",
  "Driver Contributing Circumstances + Speeding-Related (Racing / Exceeded / Too Fast)",
  "Codes 01/03/04 from SC dictionary; roadway labels from the team meeting. '≠' works on "
  "run-off-road crashes because they have one unit."),
 ("Unit Type", "Unit", "Pedalcycle, Pedestrian", "Unit Type / Non-Motorist Unit Type",
  "Numetric filter logic"),
 ("Vehicle Maneuver (unit movement)", "Unit", "Turning Left", "Vehicle Maneuver/Action",
  "Field name not confirmed - check builder"),
 ("Roadway: AADT, Through Lanes, Speed Limit, Area Type", "Roadway",
  "AADT > 20000 / <= 20000, Through Lanes = 2 / = 4 / >= 4, Speed Limit >= 35 / <= 40, Area Type = Urban / Rural",
  "Not crash-report fields (road inventory)",
  "AADT and lanes confirmed under Roadway. Speed Limit and Area Type - confirm."),
 ("Roadway Departure (derived)", "Crash",
  "Manner of Collision = Non Collision AND Crash Harmful Event not Pedestrian, Pedalcycle, Animal (deer), Animal (all other)",
  "-", "Numetric, SCDOT Custom Filter Logic, Apr 2025. Does not remove obstruction/debris - added on top."),
 ("Lane Departure (derived)", "Crash",
  "Roadway Departure OR Manner of Collision = Head On / Sideswipe, Opposite Direction",
  "-", "Numetric, Apr 2025"),
 ("Intersection Related (derived)", "Crash",
  "Junction Type = Five/More Points, Four-Way, T, Y, Traffic Circle (no driveways)",
  "-", "Numetric, Apr 2025. Runs on Junction Type, which is what production data has today. "
       "Has to be rebuilt when new-form crashes reach production."),
 ("Pedestrian Related (derived)", "Crash", "Crash Harmful Event = Pedestrian", "-", "Numetric, Apr 2025"),
]
fr.cell(1, 1, "FIELD REFERENCE").font = TITLE_FONT
for c in range(1, 6):
    fr.cell(1, c).fill = HDR_FILL
fr.row_dimensions[1].height = 25.5
for j, h in enumerate(FR_HEAD, 1):
    c = fr.cell(2, j, h)
    c.font = HDR_FONT
    c.fill = HDR_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
FW = [30, 13, 60, 42, 44]
for j, w in enumerate(FW, 1):
    fr.column_dimensions[chr(64 + j)].width = w
r = 3
for row in FR:
    for j, v in enumerate(row, 1):
        c = fr.cell(r, j, v)
        c.font = BOLD if j == 1 else BODY
        c.alignment = TOP
        c.border = BOX
    fit_height(fr, r, {3: 60, 4: 42, 5: 44})
    r += 1

r += 1
fr.cell(r, 1, "Rules for the Condition Builder").font = Font(name=F, sz=10, b=True)
r += 1
RULES = [
 "1. Group ORs in parentheses. Without them 'A AND B OR C' is read as '(A AND B) OR C', so C crashes come in "
 "without A or B (that is what happens to AADT > 20000 in a flat median barrier query).",
 "2. A list of exclusions uses AND: X ≠ a AND X ≠ b. With OR, every crash passes (a value can't equal "
 "both, so one of the '≠' is always true).",
 "3. '≠' on a unit field (Sequence of Events, Probable Cause, Unit Type, maneuver) only drops that row. "
 "The crash still comes in through its other rows. Use '≠' on crash-level fields only.",
 "4. Pick values from the dropdown. Codes here are only to check the JSON / export. The SMS codes are "
 "not the new TR-310 codes (Pedestrian is 27 in the SMS, 26 on the new form).",
 "5. New-form (MMUCC) crash reports are not in production yet. When they load, ran off road and cross "
 "median will be in Sequence of Events only, so criteria that look for them in First Harmful Event will "
 "start missing crashes.",
]
for t in RULES:
    c = fr.cell(r, 1, t)
    c.font = BODY
    c.alignment = TOP
    fr.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    fr.row_dimensions[r].height = 28
    r += 1
fr.freeze_panes = "B3"

# ---------------------------------------------------------------- Changes
ch = wb.create_sheet("Changes")
ch.cell(1, 1, "CHANGES FROM THE PREVIOUS VERSION").font = TITLE_FONT
for c in range(1, 5):
    ch.cell(1, c).fill = HDR_FILL
ch.row_dimensions[1].height = 25.5
for j, h in enumerate(["Where", "What changed", "Why", "Previous criteria"], 1):
    c = ch.cell(2, j, h)
    c.font = HDR_FONT
    c.fill = HDR_FILL
for j, w in enumerate([30, 50, 56, 60], 1):
    ch.column_dimensions[chr(64 + j)].width = w

GLOBAL = [
 ("All countermeasures",
  "Criteria rewritten with SMS field names and dropdown labels.",
  "The old column used new TR-310 fields (Relation to Junction, Type of Intersection, Alignment, "
  "Non-Motorist Unit Type, Contributing Roadway Condition) and new-form codes. Those fields aren't in the "
  "production SMS, and several codes mean something else there."),
 ("All countermeasures",
  "Codes that pointed to the wrong value in the SMS were replaced by labels.",
  "Pedestrian 26 is Pedalcycle in the SMS (Pedestrian is 27). Pedalcycle 25 is Motor Vehicle (Parked). "
  "Angle 80 doesn't exist (41-43). Tree 57 is Overhead Sign Support (Tree is 60). Utility Pole 59 is "
  "Other (Wall, Bldg) (pole is 61). Guardrail 41 is Bridge Parapet End. Median Barrier 44 is Culvert (56). "
  "Light 04-06 are shifted (SMS uses 4-7)."),
 ("Treatments, columns C-G",
  f"{len(sev_changed)} treatments: severities the study didn't cover now read 1.00.",
  "The CMF had been copied to all five severities. A KABC-only CMF was also cutting PDO crashes, and the "
  "two PDO-only CMFs were cutting fatal crashes. Column P keeps the study value; C-G are formulas."),
 ("Treatments, columns AH-AI",
  "New 'Fit with CM criteria' flag for each treatment.",
  "AASHTOWare applies the CMF only to crashes that pass the countermeasure criteria. If the CMF was "
  "measured on a different group, the benefit is off. This is the 51% vs 13% problem from the meeting."),
 ("Treatment Selector, column AA",
  "133 broken formulas (REF errors) fixed; 'Read this before you use it' column made visible again.",
  "They pointed to a Treatments column that had been deleted."),
]
r = 3
for w, what, why in GLOBAL:
    for j, v in enumerate((w, what, why), 1):
        c = ch.cell(r, j, v)
        c.font = BOLD if j == 1 else BODY
        c.alignment = TOP
        c.border = BOX
    fit_height(ch, r, {2: 50, 3: 56})
    r += 1
r += 1
ch.cell(r, 1, "By countermeasure").font = Font(name=F, sz=10, b=True)
r += 1
for src_r in range(3, 33):
    atype, loc, ctypes, ctrl, site, crit, notes, open_item = CM[src_r]
    vals = (names[src_r], notes or "Translated to SMS labels.", open_item, old_criteria[src_r])
    for j, v in enumerate(vals, 1):
        c = ch.cell(r, j, v)
        c.font = BOLD if j == 1 else BODY
        c.alignment = TOP
        c.border = BOX
    ch.cell(r, 4).font = GREY
    fit_height(ch, r, {2: 50, 3: 56, 4: 60})
    r += 1
ch.cell(2, 3).value = "Why / open item"
ch.freeze_panes = "B3"

# sheet order
order = ["Summary", "Treatment Selector", "Countermeasures", "Treatments", "Field Reference", "Changes"]
wb._sheets = [wb[n] for n in order]
wb.active = 2

wb.save(OUT)
print("saved", OUT)
print("severity rows changed:", len(sev_changed))
print("fit counts:", fit_counts)
